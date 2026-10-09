import copy
import io
import json
from unittest.mock import patch

from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from api.courses.enums import (
    AssessmentLevel,
    CourseImportStatus,
    DifficultyLevel,
    LessonContentType,
)
from api.courses.models import Assessment, CourseImportJob, LessonRequirement
from api.courses.services import course_import_template_service
from api.courses.tests.factories import (
    make_category,
    make_questions,
    make_topic,
    make_user,
)
from api.platform.services import platform_settings_service
from api.users.enums import UserRole

TEMPLATE_URL = "/api/v1/course-imports/template/"
DOWNLOAD = "api.courses.services.course_import_service.StorageService.download_bytes"


def compliant_structure() -> dict:
    """A JSON template body that meets the default submission thresholds."""

    script = "word " * 600
    return {
        "course": {
            "title": "Compliant Imported Course",
            "description": "word " * 120,
            "difficulty_level": DifficultyLevel.INTERMEDIATE,
            "tags": ["finance"],
            "learning_objectives": [f"Course objective {i}." for i in range(5)],
            "preview_video_url": "https://example.com/preview.mp4",
            "final_assessment": {"title": "Final", "questions": make_questions(15)},
        },
        "modules": [
            {
                "title": f"Module {m}",
                "order": m,
                "assessment": {"title": f"Quiz {m}", "questions": make_questions(3)},
                "lessons": [
                    {
                        "title": f"Lesson {m}-{n}",
                        "order": n,
                        "content": script,
                        "duration_minutes": 20,
                        "learning_objectives": ["Objective one.", "Objective two."],
                    }
                    for n in range(1, 4)
                ],
            }
            for m in range(1, 5)
        ],
    }


class CourseImportTemplateDownloadTests(APITestCase):
    def setUp(self):
        self.client.force_authenticate(make_user())

    def test_downloads_xlsx_template_with_live_thresholds(self):
        platform_settings_service.update_settings(course_module_count_min=5)

        response = self.client.get(TEMPLATE_URL)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response["Content-Type"], course_import_template_service.XLSX_CONTENT_TYPE
        )
        self.assertIn(
            'attachment; filename="complete-course-template.xlsx"',
            response["Content-Disposition"],
        )
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(
            workbook.sheetnames,
            ["Instructions", "Course", "Modules", "Lessons", "Questions"],
        )
        instructions = [row[0] for row in workbook["Instructions"].values if row[0]]
        self.assertTrue(any(line.startswith("Modules: 5-") for line in instructions))

    def test_downloads_json_template(self):
        response = self.client.get(TEMPLATE_URL, {"file_type": "json"})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response["Content-Type"], "application/json")
        payload = json.loads(response.content)
        self.assertIn("_instructions", payload)
        self.assertEqual(
            payload["modules"], course_import_template_service.EXAMPLE_STRUCTURE["modules"]
        )

    def test_rejects_unknown_file_type(self):
        response = self.client.get(TEMPLATE_URL, {"file_type": "pdf"})

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_requires_course_create_permission(self):
        self.client.force_authenticate(
            make_user(email="reviewer@example.com", role=UserRole.CREATOR_REVIEWER)
        )

        response = self.client.get(TEMPLATE_URL)

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)


class CourseImportTemplateImportTests(APITestCase):
    def setUp(self):
        self.creator = make_user()
        self.category = make_category()
        self.topic = make_topic(category=self.category)
        self.client.force_authenticate(self.creator)

    def _start(self, raw: bytes, file_type: str):
        content_type = course_import_template_service.TEMPLATE_FILE_TYPES[file_type]
        with patch(DOWNLOAD, return_value=raw):
            return self.client.post(
                "/api/v1/course-imports/",
                {
                    "file_key": f"uploads/course-imports/course.{file_type}",
                    "filename": f"course.{file_type}",
                    "content_type": content_type,
                    "size": len(raw),
                    "category": str(self.category.id),
                    "topic": str(self.topic.id),
                    "title": "Imported Template Course",
                    "terms_accepted": True,
                },
                format="json",
            )

    def _template(self, file_type: str) -> bytes:
        return self.client.get(TEMPLATE_URL, {"file_type": file_type}).content

    def _edited_xlsx(self, edits: dict[tuple[str, str], object]) -> bytes:
        workbook = load_workbook(io.BytesIO(self._template("xlsx")))
        for (sheet, cell), value in edits.items():
            workbook[sheet][cell] = value
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def _confirm(self, job_id):
        return self.client.post(f"/api/v1/course-imports/{job_id}/confirm/", {}, format="json")

    def test_xlsx_and_json_templates_parse_to_the_same_structure(self):
        expected, _ = course_import_template_service.validate_structure(
            copy.deepcopy(course_import_template_service.EXAMPLE_STRUCTURE)
        )

        for file_type in ("xlsx", "json"):
            with self.subTest(file_type=file_type):
                response = self._start(self._template(file_type), file_type)
                self.assertEqual(response.status_code, status.HTTP_202_ACCEPTED)
                self.assertEqual(
                    response.data["status"], CourseImportStatus.READY_FOR_REVIEW
                )
                self.assertEqual(response.data["detected_structure"], expected)
                CourseImportJob.objects.update(status=CourseImportStatus.CANCELLED)

    def test_example_template_reports_threshold_warnings(self):
        response = self._start(self._template("xlsx"), "xlsx")

        warnings = response.data["warnings"]
        self.assertIn("Course has 2 modules; submission needs 4-12.", warnings)
        self.assertIn(
            "Add a preview video in the course builder before submission.", warnings
        )

    def test_confirm_creates_the_complete_course(self):
        created = self._start(self._template("xlsx"), "xlsx")

        response = self._confirm(created.data["id"])

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn(
            "Course must have a preview video before submission (BR-015).",
            response.data["quality_failures"],
        )
        course = CourseImportJob.objects.get(pk=created.data["id"]).course
        self.assertEqual(course.title, "Personal Budgeting Essentials")
        self.assertEqual(course.difficulty_level, DifficultyLevel.BEGINNER)
        self.assertEqual(course.tags, ["budgeting", "personal finance"])
        self.assertEqual(len(course.learning_objectives), 2)
        self.assertIsNotNone(course.version_id)
        self.assertEqual(course.modules.count(), 2)
        video = course.modules.get(order=1).lessons.get(order=2)
        self.assertEqual(video.content_type, LessonContentType.VIDEO)
        self.assertEqual(video.duration_minutes, 8)
        self.assertEqual(
            video.video_url, "https://example.com/videos/budget-walkthrough.mp4"
        )
        self.assertEqual(
            list(
                LessonRequirement.objects.filter(lesson__module__course=course)
                .values_list("text", flat=True)
            ),
            ["A recent payslip or bank statement.", "A completed monthly budget."],
        )
        self.assertEqual(Assessment.objects.filter(course=course).count(), 1)
        self.assertEqual(Assessment.objects.filter(module__course=course).count(), 2)
        self.assertEqual(
            Assessment.objects.filter(lesson__module__course=course).count(), 1
        )
        module_quiz = Assessment.objects.get(
            level=AssessmentLevel.MODULE, module__course=course, module__order=1
        )
        self.assertEqual(module_quiz.questions[0]["correct_indices"], [0, 2])
        final = Assessment.objects.get(course=course)
        self.assertEqual(final.questions[1]["type"], "ESSAY")

    def test_compliant_json_template_produces_a_submittable_draft(self):
        created = self._start(json.dumps(compliant_structure()).encode(), "json")
        self.assertEqual(created.data["warnings"], [])

        response = self._confirm(created.data["id"])

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["quality_failures"], [])

    def test_xlsx_errors_point_at_the_row_to_fix(self):
        cases = [
            ({("Lessons", "A2"): 9}, "Lessons row 2: module_order 9 has no matching"),
            ({("Questions", "M2"): "x"}, "Questions row 2: correct_options must be"),
            ({("Lessons", "K2"): None}, "Lessons row 2: content - Every TEXT lesson"),
            ({("Questions", "N6"): None}, "Questions row 6: expected_answer"),
        ]
        for edits, message in cases:
            with self.subTest(message=message):
                response = self._start(self._edited_xlsx(edits), "xlsx")
                self.assertEqual(response.data["status"], CourseImportStatus.FAILED)
                self.assertIn(message, response.data["error_message"])

    def test_json_duplicate_lesson_order_fails(self):
        structure = compliant_structure()
        structure["modules"][0]["lessons"][1]["order"] = 1

        response = self._start(json.dumps(structure).encode(), "json")

        self.assertEqual(response.data["status"], CourseImportStatus.FAILED)
        self.assertIn(
            "Lesson order must be unique within each module.",
            response.data["error_message"],
        )

    def test_invalid_json_fails_with_readable_error(self):
        response = self._start(b"{not json", "json")

        self.assertEqual(response.data["status"], CourseImportStatus.FAILED)
        self.assertIn("could not be read", response.data["error_message"])
