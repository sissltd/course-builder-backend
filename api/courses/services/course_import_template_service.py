"""The downloadable complete-course import template (XLSX and JSON).

Creators who already have a finished course fill in this template and upload
it through the normal document-import flow. Both formats parse into the same
`detected_structure` tree the import review/confirm steps already use, so a
filled template carries course details, objectives, requirements and
assessments into the Draft rather than just titles and text.

The template is generated per request so its instructions always quote the
live PlatformSettings thresholds the submission quality check enforces.
"""

import copy
import io
import json
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation
from rest_framework import exceptions

from api.courses.ai.providers import GenerationStandards
from api.courses.enums import (
    AssessmentLevel,
    DifficultyLevel,
    LessonContentType,
    QuestionType,
)
from api.courses.serializers import course_import_serializer
from api.platform.services import platform_settings_service

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
JSON_CONTENT_TYPE = "application/json"
TEMPLATE_FILENAME = "complete-course-template"

MAX_MODULES = 50
MAX_LESSONS = 500
MAX_QUESTION_ROWS = 2000
MAX_OPTIONS = 6
WORDS_PER_MINUTE = 130

COURSE_SHEET = "Course"
MODULES_SHEET = "Modules"
LESSONS_SHEET = "Lessons"
QUESTIONS_SHEET = "Questions"
INSTRUCTIONS_SHEET = "Instructions"

COURSE_FIELDS = [
    "title",
    "description",
    "difficulty_level",
    "tags",
    "learning_objectives",
    "preview_video_url",
    "final_assessment_title",
]
MODULE_COLUMNS = [
    "module_order",
    "title",
    "description",
    "learning_objectives",
    "assessment_title",
]
LESSON_COLUMNS = [
    "module_order",
    "lesson_order",
    "title",
    "content_type",
    "duration_minutes",
    "learning_objectives",
    "requirements",
    "video_url",
    "embedded_link",
    "assessment_title",
    "script",
]
OPTION_COLUMNS = [f"option_{index}" for index in range(1, MAX_OPTIONS + 1)]
QUESTION_COLUMNS = [
    "scope",
    "module_order",
    "lesson_order",
    "question_order",
    "type",
    "question",
    *OPTION_COLUMNS,
    "correct_options",
    "expected_answer",
    "explanation",
    "points",
]

LIST_SEPARATOR = re.compile(r"\s*(?:\n|\|)\s*")


def _choice_question(question, options, correct_index, points=10):
    return {
        "type": QuestionType.SINGLE_CHOICE,
        "question": question,
        "options": options,
        "correct_index": correct_index,
        "points": points,
    }


#: The worked example both template formats are built from.
EXAMPLE_STRUCTURE = {
    "course": {
        "title": "Personal Budgeting Essentials",
        "description": (
            "A practical introduction to planning, tracking and improving a "
            "personal budget. Replace this with your own course description."
        ),
        "difficulty_level": DifficultyLevel.BEGINNER,
        "tags": ["budgeting", "personal finance"],
        "learning_objectives": [
            "Build a monthly budget from real income and expenses.",
            "Track spending against the budget each week.",
        ],
        "preview_video_url": "",
        "final_assessment": {
            "title": "Final Assessment",
            "questions": [
                _choice_question(
                    "What is the first step in building a budget?",
                    ["List your income", "Cut all spending", "Open a new account"],
                    0,
                ),
                {
                    "type": QuestionType.ESSAY,
                    "question": "Describe one change you will make to your budget.",
                    "expected_answer": "A specific, measurable spending change.",
                    "explanation": "Strong answers name a category and an amount.",
                    "points": 20,
                },
            ],
        },
    },
    "modules": [
        {
            "title": "Planning Your Budget",
            "description": "Set up a budget that matches how you really earn and spend.",
            "order": 1,
            "learning_objectives": ["List every regular source of income."],
            "assessment": {
                "title": "Planning Your Budget Quiz",
                "questions": [
                    {
                        "type": QuestionType.MULTIPLE_CHOICE,
                        "question": "Which of these are fixed expenses?",
                        "options": ["Rent", "Takeaway meals", "Loan repayment"],
                        "correct_indices": [0, 2],
                        "points": 10,
                    }
                ],
            },
            "lessons": [
                {
                    "title": "Knowing Your Income",
                    "order": 1,
                    "content_type": LessonContentType.TEXT,
                    "duration_minutes": 10,
                    "learning_objectives": [
                        "Identify all sources of monthly income.",
                        "Distinguish gross from net income.",
                    ],
                    "requirements": ["A recent payslip or bank statement."],
                    "video_url": "",
                    "embedded_link": "",
                    "content": "Write the full lesson script here.",
                },
                {
                    "title": "Budgeting Walkthrough",
                    "order": 2,
                    "content_type": LessonContentType.VIDEO,
                    "duration_minutes": 8,
                    "learning_objectives": [
                        "Follow a worked budget example.",
                        "Apply the 50/30/20 rule.",
                    ],
                    "requirements": [],
                    "video_url": "https://example.com/videos/budget-walkthrough.mp4",
                    "embedded_link": "",
                    "content": "Optional narration or summary for the video.",
                },
            ],
        },
        {
            "title": "Tracking and Adjusting",
            "description": "Keep the budget accurate as your month unfolds.",
            "order": 2,
            "learning_objectives": ["Review spending weekly."],
            "assessment": {
                "title": "Tracking and Adjusting Quiz",
                "questions": [
                    _choice_question(
                        "How often should you review your spending?",
                        ["Weekly", "Yearly", "Never"],
                        0,
                    )
                ],
            },
            "lessons": [
                {
                    "title": "Weekly Spending Reviews",
                    "order": 1,
                    "content_type": LessonContentType.TEXT,
                    "duration_minutes": 12,
                    "learning_objectives": [
                        "Compare actual spending to the plan.",
                        "Spot categories that overspend.",
                    ],
                    "requirements": ["A completed monthly budget."],
                    "video_url": "",
                    "embedded_link": "",
                    "content": "Write the full lesson script here.",
                },
                {
                    "title": "Check Your Understanding",
                    "order": 2,
                    "content_type": LessonContentType.QUIZ,
                    "duration_minutes": 5,
                    "learning_objectives": [
                        "Recall the weekly review steps.",
                        "Choose the right adjustment.",
                    ],
                    "requirements": [],
                    "video_url": "",
                    "embedded_link": "",
                    "content": "",
                    "assessment": {
                        "title": "Check Your Understanding",
                        "questions": [
                            _choice_question(
                                "You overspent on groceries. What do you do first?",
                                ["Find where the extra went", "Ignore it"],
                                0,
                            )
                        ],
                    },
                },
            ],
        },
    ],
}


# ---------------------------------------------------------------------------
# Template generation
# ---------------------------------------------------------------------------


def instruction_lines(standards: GenerationStandards) -> list[str]:
    """How to fill the template, quoting the live submission thresholds."""

    return [
        "How to use this template",
        "1. Replace the example rows on the Course, Modules, Lessons and Questions "
        "sheets with your own course. Keep the header rows and sheet names.",
        "2. Link rows with module_order and lesson_order (whole numbers starting at 1).",
        "3. Put several objectives, tags or requirements in one cell by separating "
        "them with a new line (Alt+Enter) or a | character.",
        "4. Questions: scope is LESSON, MODULE or COURSE. LESSON needs module_order "
        "and lesson_order, MODULE needs module_order, COURSE needs neither.",
        "5. Question type is SINGLE_CHOICE, MULTIPLE_CHOICE or ESSAY. For choice "
        "questions fill option_1 to option_6 and put the correct option number(s) "
        "in correct_options, e.g. 2 or 1,3. ESSAY questions need expected_answer "
        "and explanation instead.",
        "6. content_type is TEXT, VIDEO or QUIZ. TEXT lessons need a script; QUIZ "
        "lessons need LESSON-scope questions.",
        "7. Upload the finished file from the Import Course screen. You can review "
        "and edit everything before the draft is created.",
        "",
        "Submission requirements (your draft can be imported without meeting "
        "these, but must meet them before submission)",
        f"Course learning objectives: {standards.course_objectives_min}-"
        f"{standards.course_objectives_max}",
        f"Course description: {standards.description_words_min}-"
        f"{standards.description_words_max} words",
        f"Modules: {standards.modules_min}-{standards.modules_max}, each with a "
        "module assessment",
        f"Lessons per module: {standards.lessons_per_module_min}-"
        f"{standards.lessons_per_module_max}",
        f"Lesson learning objectives: {standards.lesson_objectives_min}-"
        f"{standards.lesson_objectives_max}",
        f"TEXT lesson script: {standards.script_words_min}-"
        f"{standards.script_words_max} words",
        f"Total course duration: {standards.duration_min_minutes}-"
        f"{standards.duration_max_minutes} minutes",
        f"Final assessment: at least {standards.final_assessment_min_questions} "
        "questions",
        "A preview video, uploaded in the course builder after import.",
    ]


TEMPLATE_FILE_TYPES = {
    "xlsx": XLSX_CONTENT_TYPE,
    "json": JSON_CONTENT_TYPE,
}


def submission_standards() -> GenerationStandards:
    """The submission quality thresholds, without AI generation's size caps."""

    return GenerationStandards.from_platform_settings(
        platform_settings_service.get_settings(), capped=False
    )


def build_template(file_type: str) -> tuple[bytes, str, str]:
    """The template file as (content, content_type, filename)."""

    standards = submission_standards()
    builder = build_template_xlsx if file_type == "xlsx" else build_template_json
    return (
        builder(standards),
        TEMPLATE_FILE_TYPES[file_type],
        f"{TEMPLATE_FILENAME}.{file_type}",
    )


def build_template_json(standards: GenerationStandards) -> bytes:
    payload = {
        "_instructions": instruction_lines(standards),
        **copy.deepcopy(EXAMPLE_STRUCTURE),
    }
    return json.dumps(payload, indent=2).encode("utf-8")


def build_template_xlsx(standards: GenerationStandards) -> bytes:
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = INSTRUCTIONS_SHEET
    instructions.column_dimensions["A"].width = 120
    for line in instruction_lines(standards):
        instructions.append([line])
    instructions["A1"].font = Font(bold=True, size=14)

    course = EXAMPLE_STRUCTURE["course"]
    course_sheet = _add_sheet(workbook, COURSE_SHEET, ["field", "value"], [30, 100])
    course_values = {
        "title": course["title"],
        "description": course["description"],
        "difficulty_level": course["difficulty_level"],
        "tags": _join(course["tags"]),
        "learning_objectives": _join(course["learning_objectives"]),
        "preview_video_url": course["preview_video_url"],
        "final_assessment_title": course["final_assessment"]["title"],
    }
    for field in COURSE_FIELDS:
        course_sheet.append([field, course_values[field]])
    _add_list_validation(course_sheet, DifficultyLevel.values, "B4")

    module_sheet = _add_sheet(
        workbook, MODULES_SHEET, MODULE_COLUMNS, [14, 40, 60, 60, 40]
    )
    lesson_sheet = _add_sheet(
        workbook,
        LESSONS_SHEET,
        LESSON_COLUMNS,
        [14, 14, 40, 14, 18, 60, 50, 40, 40, 40, 100],
    )
    question_sheet = _add_sheet(
        workbook,
        QUESTIONS_SHEET,
        QUESTION_COLUMNS,
        [12, 14, 14, 16, 20, 60, *[25] * MAX_OPTIONS, 16, 40, 40, 10],
    )
    for module in EXAMPLE_STRUCTURE["modules"]:
        module_sheet.append(
            [
                module["order"],
                module["title"],
                module["description"],
                _join(module["learning_objectives"]),
                module["assessment"]["title"],
            ]
        )
        _append_question_rows(
            question_sheet, AssessmentLevel.MODULE, module["order"], None, module["assessment"]
        )
        for lesson in module["lessons"]:
            assessment = lesson.get("assessment")
            lesson_sheet.append(
                [
                    module["order"],
                    lesson["order"],
                    lesson["title"],
                    lesson["content_type"],
                    lesson["duration_minutes"],
                    _join(lesson["learning_objectives"]),
                    _join(lesson["requirements"]),
                    lesson["video_url"],
                    lesson["embedded_link"],
                    assessment["title"] if assessment else "",
                    lesson["content"],
                ]
            )
            if assessment:
                _append_question_rows(
                    question_sheet,
                    AssessmentLevel.LESSON,
                    module["order"],
                    lesson["order"],
                    assessment,
                )
    _append_question_rows(
        question_sheet, AssessmentLevel.COURSE, None, None, course["final_assessment"]
    )
    _add_list_validation(lesson_sheet, LessonContentType.values, "D2:D1000")
    _add_list_validation(question_sheet, AssessmentLevel.values, "A2:A5000")
    _add_list_validation(question_sheet, QuestionType.values, "E2:E5000")

    for sheet in (course_sheet, module_sheet, lesson_sheet, question_sheet):
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _add_sheet(workbook, title, headers, widths):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell, width in zip(sheet[1], widths, strict=True):
        cell.font = Font(bold=True)
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.freeze_panes = "A2"
    return sheet


def _add_list_validation(sheet, values, cell_range):
    validation = DataValidation(
        type="list", formula1=f'"{",".join(values)}"', allow_blank=True
    )
    sheet.add_data_validation(validation)
    validation.add(cell_range)


def _append_question_rows(sheet, scope, module_order, lesson_order, assessment):
    for question_order, question in enumerate(assessment["questions"], 1):
        options = question.get("options") or []
        if "correct_indices" in question:
            correct = ",".join(str(index + 1) for index in question["correct_indices"])
        elif "correct_index" in question:
            correct = str(question["correct_index"] + 1)
        else:
            correct = ""
        sheet.append(
            [
                scope,
                module_order,
                lesson_order,
                question_order,
                question["type"],
                question["question"],
                *(options + [""] * (MAX_OPTIONS - len(options))),
                correct,
                question.get("expected_answer", ""),
                question.get("explanation", ""),
                question.get("points", 0),
            ]
        )


def _join(items: list[str]) -> str:
    return "\n".join(items)


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------


def parse_template_json(raw: bytes) -> tuple[dict, list[str]]:
    try:
        payload = json.loads(raw.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise exceptions.ValidationError(
            "This JSON file could not be read. Check it is valid JSON and retry."
        ) from exc
    if not isinstance(payload, dict):
        raise exceptions.ValidationError(
            "The JSON template must be an object with 'course' and 'modules' keys."
        )
    payload.pop("_instructions", None)
    modules = payload.get("modules")
    if isinstance(modules, list):
        _check_caps(
            module_count=len(modules),
            lesson_count=sum(
                len(module.get("lessons") or [])
                for module in modules
                if isinstance(module, dict)
            ),
        )
    return validate_structure(payload)


def parse_template_xlsx(raw: bytes) -> tuple[dict, list[str]]:
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise exceptions.ValidationError(
            "This XLSX file could not be opened. Re-save it as an Excel "
            "workbook (.xlsx) and retry."
        ) from exc
    try:
        for name in (COURSE_SHEET, MODULES_SHEET, LESSONS_SHEET):
            if name not in workbook.sheetnames:
                raise exceptions.ValidationError(
                    f"The workbook is missing the '{name}' sheet. Download a "
                    "fresh template and copy your content into it."
                )
        course_rows = _sheet_rows(workbook[COURSE_SHEET])
        module_rows = _sheet_rows(workbook[MODULES_SHEET])
        lesson_rows = _sheet_rows(workbook[LESSONS_SHEET])
        question_rows = (
            _sheet_rows(workbook[QUESTIONS_SHEET])
            if QUESTIONS_SHEET in workbook.sheetnames
            else []
        )
    finally:
        workbook.close()
    _check_caps(module_count=len(module_rows), lesson_count=len(lesson_rows))
    if len(question_rows) > MAX_QUESTION_ROWS:
        raise exceptions.ValidationError(
            f"The template can hold at most {MAX_QUESTION_ROWS} question rows."
        )
    structure, locations = _assemble_structure(
        course_rows, module_rows, lesson_rows, question_rows
    )
    return validate_structure(structure, locations=locations)


def validate_structure(
    structure: dict, *, locations: dict[str, str] | None = None
) -> tuple[dict, list[str]]:
    """Validate an import tree and return it normalized, plus threshold warnings.

    `locations` maps structure paths (e.g. ``modules[0].lessons[1]``) to
    spreadsheet row labels so errors point at the row the creator must fix.
    """

    serializer = course_import_serializer.ImportedStructureSerializer(data=structure)
    if not serializer.is_valid():
        messages = _flatten_errors(serializer.errors)
        raise exceptions.ValidationError(
            " ".join(_locate(path, message, locations or {}) for path, message in messages)
        )
    validated = json.loads(json.dumps(serializer.validated_data))
    return validated, threshold_warnings(validated)


def threshold_warnings(
    structure: dict, standards: GenerationStandards | None = None
) -> list[str]:
    """Non-blocking notes on where the tree falls short of submission standards."""

    standards = standards or submission_standards()
    warnings: list[str] = []
    course = structure.get("course") or {}
    modules = structure["modules"]

    def out_of_range(value, low, high):
        return not low <= value <= high

    objective_count = len(course.get("learning_objectives") or [])
    if out_of_range(
        objective_count, standards.course_objectives_min, standards.course_objectives_max
    ):
        warnings.append(
            f"Course has {objective_count} learning objectives; submission needs "
            f"{standards.course_objectives_min}-{standards.course_objectives_max}."
        )
    if "description" in course:
        description_words = len(course["description"].split())
        if out_of_range(
            description_words,
            standards.description_words_min,
            standards.description_words_max,
        ):
            warnings.append(
                f"Course description has {description_words} words; submission "
                f"needs {standards.description_words_min}-"
                f"{standards.description_words_max}."
            )
    if out_of_range(len(modules), standards.modules_min, standards.modules_max):
        warnings.append(
            f"Course has {len(modules)} modules; submission needs "
            f"{standards.modules_min}-{standards.modules_max}."
        )
    total_minutes = 0
    for module in modules:
        lessons = module["lessons"]
        if out_of_range(
            len(lessons), standards.lessons_per_module_min, standards.lessons_per_module_max
        ):
            warnings.append(
                f"Module '{module['title']}' has {len(lessons)} lessons; submission "
                f"needs {standards.lessons_per_module_min}-"
                f"{standards.lessons_per_module_max}."
            )
        if not module.get("assessment"):
            warnings.append(f"Module '{module['title']}' has no module assessment.")
        for lesson in lessons:
            total_minutes += lesson_duration_minutes(lesson)
            lesson_objectives = len(lesson.get("learning_objectives") or [])
            if out_of_range(
                lesson_objectives,
                standards.lesson_objectives_min,
                standards.lesson_objectives_max,
            ):
                warnings.append(
                    f"Lesson '{lesson['title']}' has {lesson_objectives} learning "
                    f"objectives; submission needs {standards.lesson_objectives_min}-"
                    f"{standards.lesson_objectives_max}."
                )
            if lesson.get("content_type", LessonContentType.TEXT) == LessonContentType.TEXT:
                script_words = len(lesson["content"].split())
                if out_of_range(
                    script_words, standards.script_words_min, standards.script_words_max
                ):
                    warnings.append(
                        f"Lesson '{lesson['title']}' script has {script_words} words; "
                        f"submission needs {standards.script_words_min}-"
                        f"{standards.script_words_max}."
                    )
    if out_of_range(
        total_minutes, standards.duration_min_minutes, standards.duration_max_minutes
    ):
        warnings.append(
            f"Course runs {total_minutes} minutes; submission needs "
            f"{standards.duration_min_minutes}-{standards.duration_max_minutes}."
        )
    final_questions = len((course.get("final_assessment") or {}).get("questions") or [])
    if final_questions < standards.final_assessment_min_questions:
        warnings.append(
            f"Final assessment has {final_questions} questions; submission needs at "
            f"least {standards.final_assessment_min_questions}."
        )
    if not course.get("preview_video_url"):
        warnings.append("Add a preview video in the course builder before submission.")
    return warnings


def lesson_duration_minutes(lesson: dict) -> int:
    """The lesson's stated duration, or an estimate from its script length."""

    return lesson.get("duration_minutes") or max(
        1, len((lesson.get("content") or "").split()) // WORDS_PER_MINUTE
    )


def _check_caps(*, module_count: int, lesson_count: int) -> None:
    if module_count > MAX_MODULES:
        raise exceptions.ValidationError(
            f"The template can hold at most {MAX_MODULES} modules."
        )
    if lesson_count > MAX_LESSONS:
        raise exceptions.ValidationError(
            f"The template can hold at most {MAX_LESSONS} lessons."
        )


def _sheet_rows(sheet) -> list[tuple[int, dict]]:
    """Non-blank rows as (row_number, {normalized_header: value})."""

    rows = sheet.iter_rows(values_only=True)
    header_row = next(rows, None) or ()
    headers = [
        str(header).strip().lower().replace(" ", "_") if header is not None else ""
        for header in header_row
    ]
    parsed = []
    for row_number, values in enumerate(rows, 2):
        record = {
            header: value
            for header, value in zip(headers, values)
            if header and value is not None and str(value).strip() != ""
        }
        if record:
            parsed.append((row_number, record))
    return parsed


def _text(record: dict, key: str) -> str:
    value = record.get(key)
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _number(record: dict, key: str, label: str) -> int | None:
    value = _text(record, key)
    if not value:
        return None
    try:
        return int(float(value))
    except ValueError:
        raise exceptions.ValidationError(
            f"{label}: {key} must be a whole number (got '{value}')."
        ) from None


def _items(record: dict, key: str) -> list[str]:
    return [item for item in LIST_SEPARATOR.split(_text(record, key)) if item]


def _assemble_structure(course_rows, module_rows, lesson_rows, question_rows):
    locations: dict[str, str] = {}
    course_values = {
        _text(record, "field").lower(): record.get("value")
        for _, record in course_rows
    }
    course_record = {key: value for key, value in course_values.items() if value is not None}
    course: dict = {}
    for field in ("title", "description", "difficulty_level"):
        if _text(course_record, field):
            course[field] = _text(course_record, field)
    course["preview_video_url"] = _text(course_record, "preview_video_url")
    for field in ("tags", "learning_objectives"):
        if field in course_record:
            course[field] = _items(course_record, field)

    modules: list[dict] = []
    module_index_by_order: dict[int, int] = {}
    for row_number, record in module_rows:
        label = f"{MODULES_SHEET} row {row_number}"
        order = _number(record, "module_order", label)
        if order is None:
            raise exceptions.ValidationError(f"{label}: module_order is required.")
        if order in module_index_by_order:
            raise exceptions.ValidationError(
                f"{label}: module_order {order} is used more than once."
            )
        module_index_by_order[order] = len(modules)
        locations[f"modules[{len(modules)}]"] = label
        modules.append(
            {
                "title": _text(record, "title"),
                "description": _text(record, "description"),
                "order": order,
                "learning_objectives": _items(record, "learning_objectives"),
                "lessons": [],
                "_assessment_title": _text(record, "assessment_title"),
            }
        )

    lesson_index_by_key: dict[tuple[int, int], tuple[int, int]] = {}
    for row_number, record in lesson_rows:
        label = f"{LESSONS_SHEET} row {row_number}"
        module_order = _number(record, "module_order", label)
        if module_order not in module_index_by_order:
            raise exceptions.ValidationError(
                f"{label}: module_order {module_order} has no matching row on the "
                f"{MODULES_SHEET} sheet."
            )
        module_index = module_index_by_order[module_order]
        lessons = modules[module_index]["lessons"]
        lesson_order = _number(record, "lesson_order", label)
        if lesson_order is not None:
            lesson_index_by_key[(module_order, lesson_order)] = (
                module_index,
                len(lessons),
            )
        locations[f"modules[{module_index}].lessons[{len(lessons)}]"] = label
        lesson = {
            "title": _text(record, "title"),
            "content": _text(record, "script"),
            "content_type": _text(record, "content_type").upper() or LessonContentType.TEXT,
            "learning_objectives": _items(record, "learning_objectives"),
            "requirements": _items(record, "requirements"),
            "video_url": _text(record, "video_url"),
            "embedded_link": _text(record, "embedded_link"),
            "_assessment_title": _text(record, "assessment_title"),
        }
        if lesson_order is not None:
            lesson["order"] = lesson_order
        duration = _number(record, "duration_minutes", label)
        if duration is not None:
            lesson["duration_minutes"] = duration
        lessons.append(lesson)

    assessments: dict[tuple, list[tuple[int, int, dict]]] = {}
    for row_number, record in question_rows:
        label = f"{QUESTIONS_SHEET} row {row_number}"
        scope = _text(record, "scope").upper()
        module_order = _number(record, "module_order", label)
        lesson_order = _number(record, "lesson_order", label)
        if scope == AssessmentLevel.COURSE:
            key = (AssessmentLevel.COURSE,)
        elif scope == AssessmentLevel.MODULE:
            if module_order not in module_index_by_order:
                raise exceptions.ValidationError(
                    f"{label}: module_order {module_order} has no matching module."
                )
            key = (AssessmentLevel.MODULE, module_index_by_order[module_order])
        elif scope == AssessmentLevel.LESSON:
            if (module_order, lesson_order) not in lesson_index_by_key:
                raise exceptions.ValidationError(
                    f"{label}: no lesson matches module_order {module_order} and "
                    f"lesson_order {lesson_order}."
                )
            key = (AssessmentLevel.LESSON, *lesson_index_by_key[(module_order, lesson_order)])
        else:
            raise exceptions.ValidationError(
                f"{label}: scope must be LESSON, MODULE, or COURSE."
            )
        question_order = _number(record, "question_order", label) or row_number
        assessments.setdefault(key, []).append(
            (question_order, row_number, _question_from_row(record, label))
        )

    for key, rows in assessments.items():
        rows.sort(key=lambda item: (item[0], item[1]))
        questions = [question for _, _, question in rows]
        if key[0] == AssessmentLevel.COURSE:
            path = "course.final_assessment"
            course["final_assessment"] = {
                "title": _text(course_record, "final_assessment_title") or "Final Assessment",
                "questions": questions,
            }
        elif key[0] == AssessmentLevel.MODULE:
            module = modules[key[1]]
            path = f"modules[{key[1]}].assessment"
            module["assessment"] = {
                "title": module["_assessment_title"] or f"{module['title']} Quiz",
                "questions": questions,
            }
        else:
            lesson = modules[key[1]]["lessons"][key[2]]
            path = f"modules[{key[1]}].lessons[{key[2]}].assessment"
            lesson["assessment"] = {
                "title": lesson["_assessment_title"] or lesson["title"],
                "questions": questions,
            }
        for index, (_, row_number, _) in enumerate(rows):
            locations[f"{path}.questions[{index}]"] = f"{QUESTIONS_SHEET} row {row_number}"

    for module in modules:
        module.pop("_assessment_title")
        for lesson in module["lessons"]:
            lesson.pop("_assessment_title")
    structure: dict = {"modules": modules}
    if course:
        structure["course"] = course
    return structure, locations


def _question_from_row(record: dict, label: str) -> dict:
    question_type = _text(record, "type").upper() or QuestionType.SINGLE_CHOICE
    question = {"type": question_type, "question": _text(record, "question")}
    points = _number(record, "points", label)
    if points is not None:
        question["points"] = points
    for field in ("expected_answer", "explanation"):
        if _text(record, field):
            question[field] = _text(record, field)
    if question_type == QuestionType.ESSAY:
        return question
    question["options"] = [
        _text(record, column) for column in OPTION_COLUMNS if _text(record, column)
    ]
    raw_correct = _text(record, "correct_options")
    try:
        correct = [int(part) - 1 for part in re.split(r"[,\s;]+", raw_correct) if part]
    except ValueError:
        raise exceptions.ValidationError(
            f"{label}: correct_options must be option numbers like 2 or 1,3 "
            f"(got '{raw_correct}')."
        ) from None
    if not correct:
        raise exceptions.ValidationError(f"{label}: correct_options is required.")
    if question_type == QuestionType.SINGLE_CHOICE:
        if len(correct) != 1:
            raise exceptions.ValidationError(
                f"{label}: SINGLE_CHOICE questions have exactly one correct option."
            )
        question["correct_index"] = correct[0]
    else:
        question["correct_indices"] = correct
    return question


def _flatten_errors(errors, path: str = "") -> list[tuple[str, str]]:
    if isinstance(errors, dict):
        flattened = []
        for key, value in errors.items():
            if key == "non_field_errors":
                flattened.extend(_flatten_errors(value, path))
            else:
                flattened.extend(_flatten_errors(value, f"{path}.{key}" if path else key))
        return flattened
    if isinstance(errors, list):
        if all(isinstance(item, str) for item in errors):
            return [(path, str(item)) for item in errors]
        flattened = []
        for index, item in enumerate(errors):
            if item:
                flattened.extend(_flatten_errors(item, f"{path}[{index}]"))
        return flattened
    return [(path, str(errors))]


def _locate(path: str, message: str, locations: dict[str, str]) -> str:
    """Prefix an error with its spreadsheet row, or its structure path."""

    for prefix in sorted(locations, key=len, reverse=True):
        if path == prefix or path.startswith(f"{prefix}."):
            field = path[len(prefix) :].lstrip(".")
            return f"{locations[prefix]}: {field + ' - ' if field else ''}{message}"
    return f"{path}: {message}" if path else message
